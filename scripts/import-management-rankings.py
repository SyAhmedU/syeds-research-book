"""Read Scopus Sources XLSX exports verbatim into reproducible ranking evidence.

Does not edit workbooks or infer ISSNs, source IDs, or management-category ranks.
Run the journal builder after this script to apply unique exact-title matches.
"""
import argparse
import hashlib
import json
import re
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument('inputs', nargs='+', type=Path)
args = parser.parse_args()
records, files = [], []
years = set()
for file in args.inputs:
    workbook = openpyxl.load_workbook(file, read_only=True, data_only=True)
    file_rows = 0
    for sheet in workbook:
        rows = sheet.iter_rows(values_only=True)
        headers = list(next(rows))
        required = {'Source title', 'CiteScore', 'Highest percentile', 'Publisher'}
        if not required.issubset(headers):
            raise ValueError(f'{file.name}/{sheet.title}: missing required columns')
        citation_headers = [h for h in headers if isinstance(h, str) and re.fullmatch(r'20\d{2}-\d{2} Citations', h)]
        if len(citation_headers) != 1:
            raise ValueError(f'{file.name}: cannot determine citation window')
        citation_header = citation_headers[0]
        start, end = re.match(r'(20\d{2})-(\d{2})', citation_header).groups()
        year = int(start[:2] + end)
        if year - int(start) != 3:
            raise ValueError('Expected a four-year CiteScore window')
        years.add(year)
        for row_number, values in enumerate(rows, 2):
            if not any(v is not None for v in values):
                continue
            raw = dict(zip(headers, values))
            name = raw['Source title']
            text = raw['Highest percentile']
            match = re.fullmatch(r'(\d+(?:\.\d+)?)%\s+(\d+)/(\d+)\s*(.*)', str(text), re.DOTALL)
            if not isinstance(name, str) or not match:
                raise ValueError(f'{file.name} row {row_number}: invalid source/percentile')
            percentile, rank, total, category = match.groups()
            percentile, rank, total = float(percentile), int(rank), int(total)
            if not 0 <= percentile <= 99 or not 1 <= rank <= total:
                raise ValueError(f'{file.name} row {row_number}: invalid percentile/rank range')
            quartile = 'Q1' if percentile >= 75 else 'Q2' if percentile >= 50 else 'Q3' if percentile >= 25 else 'Q4'
            records.append({
                'name': name, 'citeScore': raw['CiteScore'], 'highestPercentile': percentile,
                'highestPercentileRaw': text, 'rank': rank, 'categoryTotal': total,
                'highestPercentileCategory': category.strip() or None, 'quartile': quartile,
                'quartileScope': 'highest-across-categories', 'rankingYear': year,
                'rankingYearBasis': f'four-year citation window in export header: {citation_header}',
                'publisher': raw['Publisher'], 'raw': raw,
                'evidence': {'file': file.name, 'sheet': sheet.title, 'row': row_number,
                             'percentileCell': f'{openpyxl.utils.get_column_letter(headers.index("Highest percentile") + 1)}{row_number}'},
            })
            file_rows += 1
    workbook.close()
    files.append({'name': file.name, 'sha256': hashlib.sha256(file.read_bytes()).hexdigest(), 'records': file_rows})
if len(years) != 1:
    raise ValueError('Mixed ranking years: import separate snapshots instead')
output = {
    'version': 1, 'importedAt': datetime.now(timezone.utc).isoformat(),
    'source': 'User-supplied Scopus Sources exports', 'rankingYear': years.pop(),
    'quartileScope': 'highest-across-categories',
    'quartileRule': 'Q1: 75–99; Q2: 50–74; Q3: 25–49; Q4: 0–24; derived from exported highest percentile',
    'quartileRuleUrl': 'https://www.elsevier.support/scopus/answer/how-do-i-use-the-scopus-sources-feature',
    'identityNote': 'Exports contain no ISSN or Scopus source ID. Journal attachment is a unique exact normalized title match, not a verified identifier match.',
    'files': files, 'coverage': {'records': len(records), 'quartiles': dict(Counter(r['quartile'] for r in records))},
    'records': records,
}
target = ROOT / 'data' / 'management-rankings.json'
temporary = target.with_suffix('.json.tmp')
temporary.write_text(json.dumps(output, ensure_ascii=False, separators=(',', ':')), encoding='utf-8')
temporary.replace(target)
print(json.dumps(output['coverage']))
